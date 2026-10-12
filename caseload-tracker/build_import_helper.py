"""Build the import helper: rearranges an old caseload sheet into the tracker's columns.

Usage:
    python build_import_helper.py [output.xlsx] [--sample] [--cardex]

The old sheet is pasted as-is onto 'Paste old sheet'. 'Converted' finds each old column by its
heading (in the row named on 'Start here', C3) and lays the data out in the tracker's column
order, ready to paste as values into the Caseload tab. --cardex builds the converter for the
cardex-style sheet (headings in row 2, GP / Delivering at / Complications / 41/40).
--sample fills a fictional old sheet for testing.
"""
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import build_tracker as bt

ROWS = 400
DATA0 = 6                      # first data row on 'Converted'
PASTE = "'Paste old sheet'"
PASTE_RANGE = f"{PASTE}!$A$1:$BZ${ROWS + 50}"
HDR_ROW = "'Start here'!$C$3"   # row of the old sheet's column headings (1-10)
# The heading row as a range: CHOOSE over fixed rows keeps this to plain, non-volatile Excel.
HEADINGS = f"CHOOSE({HDR_ROW}," + ",".join(f"{PASTE}!$A${i}:$BZ${i}" for i in range(1, 11)) + ")"


def old_headers():
    """Tracker column -> start of the matching heading in the old sheet (case-insensitive)."""
    m = {"Name": "name", "NHS number": "nhs", "DOB": "dob", "Parity": "parity", "MLC/CLC": "mlc",
         "EDD": "edd", "Booking date": "booking", "Risks": "info", **{a: a for a in bt.APPTS}}
    if bt.VARIANT == "cardex":
        # 'deliver' also matches the cardex's 'Deliverying at' spelling
        m.update({"GP": "gp", "Delivering at": "deliver", "Risks": "complications"})
        del m["Booking date"]
    return m


def auto_cols():
    return {c for c, _, g in bt.COLUMNS if g == bt.G_AUTO}


def blocks():
    """Runs of consecutive tracker input columns: [(first_header, last_header), ...]."""
    runs, cur = [], []
    for h, _, _ in bt.COLUMNS:
        if h in auto_cols():
            if cur:
                runs.append((cur[0], cur[-1]))
            cur = []
        else:
            cur.append(h)
    if cur:
        runs.append((cur[0], cur[-1]))
    return runs


def build(path, sample=False):
    wb = Workbook()
    font = Font(name=bt.FONT, size=10)
    runs = blocks()
    AUTO, OLD = auto_cols(), old_headers()
    DATE_COLS = {"DOB", "EDD", *bt.CHAIN, *bt.DATE_ACTIONS}
    cardex = bt.VARIANT == "cardex"

    # ---- Start here -------------------------------------------------------
    st = wb.active
    st.title = "Start here"
    st.column_dimensions["A"].width = 3
    st.column_dimensions["B"].width = 105
    st["B2"] = "Moving the old caseload sheet into the new tracker"
    st["B2"].font = Font(name=bt.FONT, size=14, bold=True, color="1F4E79")
    st["B3"] = "Row number of the column headings in the pasted sheet:"
    st["B3"].alignment = Alignment(horizontal="right")
    st["B3"].font = Font(name=bt.FONT, size=10, bold=True)
    st["C3"] = 2 if cardex else 1
    st["C3"].fill, st["C3"].font = bt.INPUT_YELLOW, Font(name=bt.FONT, size=12, bold=True, color="0000FF")
    st["C3"].alignment = Alignment(horizontal="center")
    st.column_dimensions["C"].width = 8
    v = DataValidation(type="whole", operator="between", formula1="1", formula2="10",
                       showErrorMessage=True, errorTitle="Heading row", error="Enter a row number from 1 to 10.")
    v.add("C3")
    st.add_data_validation(v)
    steps = [
        ("", "Everything happens inside Excel on your work computer. Patient data never needs to be "
             "emailed or uploaded anywhere."),
        ("s", "1. Paste the old sheet"),
        ("", "Open the old sheet, click the grey box top-left of the grid (selects everything) and Ctrl+C. "
             "Come back here, click cell A1 on the 'Paste old sheet' tab and Ctrl+V."),
        ("", "Check the yellow box above (C3) matches the row the column headings (Name, NHS, EDD, 16/40 ...) "
             "are in on the 'Paste old sheet' tab. " + ("It's set to 2 for the cardex, whose headings are in "
             "row 2." if cardex else "It's set to 1.") + " Everything below the headings is treated as patients."),
        ("s", "2. Check the matching"),
        ("", "On the 'Converted' tab, row 4 shows which old heading was found for each new column."),
        ("", "• 'NOT FOUND' (red): the old heading is spelled differently. Change the yellow cell in row 3 "
             "to the start of the old heading (capitals don't matter)."),
        ("", "• '(new column)': the old sheet didn't have it (" + ("Status, Booking date, GTT, Anti-D, IOL ..."
             if cardex else "GTT, Anti-D, IOL ...") + "), so it's left blank."),
        ("", "• Orange cells in the data: a date column holds text Excel can't read as a date "
             "(e.g. '12/3' or 'seen'). Fix these in the tracker after pasting."),
        ("s", "3. Copy into the tracker (paste as VALUES)"),
    ]
    for i, (a, b) in enumerate(runs, start=1):
        rng = f"{bt.COL[a]}{DATA0}:{bt.COL[b]}{DATA0 + ROWS - 1}"
        steps.append(("", f"Block {i}: on 'Converted' select {rng} ({a} to {b}), Ctrl+C. In the tracker's "
                          f"Caseload tab click {bt.COL[a]}{bt.FIRST}, then Home > Paste > Paste Values "
                          f"(the clipboard icon with 123). Or right-click > Paste Options > Values."))
    steps += [
        ("", "Paste Values keeps the tracker's formatting and doesn't overwrite the grey calculated "
             "columns, which is why there are separate blocks."),
        ("s", "4. Tidy up in the tracker"),
        ("", "• Set Status for anyone who has delivered or left, so their rows grey out."),
        ("", "• Past appointment dates with nothing after them will show red. That's the tracker doing "
             "its job: check each one."),
        ("", "• Fill in the new columns (" + ("Booking date, " if cardex else "") + "GTT, Anti-D, IOL, BP, "
             "Last updated by) as you go. The old " + ("Complications" if cardex else "info") +
             " column goes into Risks."),
        ("s", "5. Delete this helper file"),
        ("", "It now holds a copy of patient data. Delete it once the tracker looks right."),
    ]
    for i, (kind, text) in enumerate(steps, start=5):
        c = st.cell(row=i, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = (Font(name=bt.FONT, size=14, bold=True, color="1F4E79") if kind == "h" else
                  Font(name=bt.FONT, size=11, bold=True, color="1F4E79") if kind == "s" else font)

    # ---- Paste old sheet --------------------------------------------------
    ps = wb.create_sheet("Paste old sheet")

    # ---- Converted --------------------------------------------------------
    cv = wb.create_sheet("Converted")
    block_of = {}
    for i, (a, b) in enumerate(runs, start=1):
        first, last = cv[f"{bt.COL[a]}1"].column, cv[f"{bt.COL[b]}1"].column
        for c in range(first, last + 1):
            block_of[c] = i
    for idx, (h, w, _) in enumerate(bt.COLUMNS, start=1):
        L = get_column_letter(idx)
        cv.column_dimensions[L].width = max(w, 12)
        skip = h in AUTO
        b1 = cv[f"{L}1"]
        b1.value = "calculated: skip" if skip else f"BLOCK {block_of[idx]}"
        b1.fill = bt.GREY if skip else bt.fill("DDEBF7" if block_of[idx] % 2 else "E2EFDA")
        b1.font = Font(name=bt.FONT, size=9, bold=True, color="595959")
        hdr = cv[f"{L}2"]
        hdr.value, hdr.fill = h, bt.HEAD_FILL
        hdr.font = Font(name=bt.FONT, size=10, bold=True, color="FFFFFF")
        hdr.alignment = Alignment(horizontal="center", wrap_text=True)
        if skip:
            for r in range(3, 6):
                cv[f"{L}{r}"].fill = bt.GREY
            continue
        look = cv[f"{L}3"]
        look.value, look.fill = OLD.get(h, ""), bt.INPUT_YELLOW
        look.font = Font(name=bt.FONT, size=9, color="0000FF")
        cv[f"{L}5"] = f'=IF({L}3="","",IFERROR(MATCH({L}3&"*",{HEADINGS},0),""))'
        cv[f"{L}4"] = (f'=IF({L}3="","(new column)",IF(ISNUMBER({L}5),INDEX({HEADINGS},{L}5),'
                       f'"NOT FOUND"))')
        cv[f"{L}4"].font = Font(name=bt.FONT, size=9, italic=True)
        cv[f"{L}5"].font = Font(name=bt.FONT, size=8, color="A6A6A6")
        get = f"INDEX({PASTE_RANGE},{HDR_ROW}+ROW()-{DATA0 - 1},{L}$5)"
        for r in range(DATA0, DATA0 + ROWS):
            cv[f"{L}{r}"] = f'=IF(ISNUMBER({L}$5),IF({get}="","",{get}),"")'
            cv[f"{L}{r}"].font = font
            if h in DATE_COLS:
                cv[f"{L}{r}"].number_format = bt.DATE_FMT
        if h in DATE_COLS:
            x = f"{L}{DATA0}"
            cv.conditional_formatting.add(
                f"{L}{DATA0}:{L}{DATA0 + ROWS - 1}",
                FormulaRule(formula=[f'AND(ISTEXT({x}),{x}<>"",{x}<>"DNA",{x}<>"N/A")'],
                            fill=bt.ORANGE, font=Font(color="833C0B")))
    last_col = get_column_letter(len(bt.COLUMNS))
    cv.conditional_formatting.add(f"A4:{last_col}4", FormulaRule(
        formula=['A4="NOT FOUND"'], fill=bt.RED, font=Font(color=bt.RED_TXT, bold=True)))
    cv.freeze_panes = f"B{DATA0}"
    cv.row_dimensions[2].height = 30

    if sample:
        (add_sample_cardex if cardex else add_sample)(ps)
    wb.active = 0
    wb.save(path)


def add_sample(ps):
    """Fictional old sheet in the old column order, with messy headers and a text date."""
    t = date.today()
    headers = ["Name", "NHS number", "dob", "parity", "mlc/clc", "edd", "booking date",
               "16/40 ", "25/40 ", "28/40 ", "31/40 ", "34/40 ", "36/40 ", "38/40 ", "40/40 ", "info"]
    ps.append(headers)
    ps.append(["Test One", 9993982598, date(1995, 3, 14), "G1P0", "MLC", t + timedelta(days=70),
               t - timedelta(days=150), t - timedelta(days=100), t - timedelta(days=35), t + timedelta(days=5),
               None, None, None, None, None, "note one"])
    ps.append(["Test Two", "999 791 9076", date(1990, 7, 2), "G2P1", "CLC", t + timedelta(days=20),
               t - timedelta(days=180), t - timedelta(days=150), "N/A", "12/3", "DNA", None, None, None, None, ""])


def add_sample_cardex(ps):
    """Fictional cardex: blank column A, headings in row 2 (with its spellings), 41/40."""
    t = date.today()
    d = lambda n: t + timedelta(days=n)
    ps.append([])
    ps.append([None, "Name", "NHS", "DOB", "Parity", "GP", "MLC/CLC", "Deliverying at", "EDD", "Complications",
               "16/40", "25/40", "28/40", "31/40", "34/40", "36/40", "38/40", "40/40", "41/40"])
    ps.append([None, "Test Cardex", 9993982598, date(1995, 3, 14), "G1P0", "Dr Test", "CLC", "Test Hospital",
               d(-5), "GDM", d(-170), d(-100), d(-79), d(-58), d(-37), d(-23), d(-9), d(-5), d(2)])
    ps.append([None, "Test Two", "999 791 9076", date(1990, 7, 2), "G2P1", "Test Surgery", "MLC", "Home",
               d(40), "", d(-100), "N/A", "12/3", None, None, None, None, None, None])


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if "--cardex" in sys.argv:
        bt.configure("cardex")
    build(args[0] if args else "Import_Helper.xlsx", sample="--sample" in sys.argv)
